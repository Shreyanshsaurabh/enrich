"""
Developer pipeline:
  1. Keep only apps that contain ads
  2. Deduplicate by Developer ID
       - Rating     -> average across the developer's apps
       - Downloads  -> sum across the developer's apps
       - Genre, Category, App Name, App Link, Privacy Policy -> from the developer's
         most-downloaded app
  3. Visit each developer's privacy policy page and extract privacy_email / privacy_number
  4. Write the result to an Excel file

Install:
    pip install pandas openpyxl requests beautifulsoup4 phonenumbers
    pip install playwright && python -m playwright install chromium     # only for --browser-fallback

First full run:
    python developer_pipeline.py all_apps.xlsx out.xlsx

Second pass, ONLY on rows that failed (rows that already worked are never fetched again):
    python developer_pipeline.py all_apps.xlsx out.xlsx --retry-failed --browser-fallback

  --retry-failed      re-fetch rows that failed for transient reasons (429, 5xx, timeouts,
                      connection errors, crashes) using exponential backoff
  --browser-fallback  re-try rows with no_contact_found / 403 / 406 / 412 / 429 in a real
                      headless browser (runs JavaScript) and follow the site's "Contact" link
  --max-minutes N     stop gracefully after N minutes, still write the Excel (exit code 3);
                      re-run the same command to continue

Scraping results are cached in <output>.cache.json (keyed by privacy URL), so every run
resumes from where the previous one stopped.
"""
import argparse
import asyncio
import json
import os
import random
import re
import sys
import threading
import time
import traceback
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import urljoin, urlparse

import pandas as pd
import phonenumbers
import requests
from bs4 import BeautifulSoup
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; PrivacyContactBot/1.0)"}
BROWSER_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36")
TIMEOUT = 15
MAX_BYTES = 2_000_000
HOST_DELAY = 1.0          # min seconds between requests to the same host (--host-delay)
DEADLINE = None           # set by --max-minutes
MAX_EMAILS = 3
MAX_PHONES = 3

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9-]+(?:\.[a-zA-Z0-9-]+)*\.[a-zA-Z]{2,}")
OBFUSCATED_RE = re.compile(
    r"([a-zA-Z0-9._%+-]+)\s*[\[\(]\s*(?:at|AT)\s*[\]\)]\s*"
    r"([a-zA-Z0-9-]+(?:\s*[\[\(]?\s*(?:dot|DOT)\s*[\]\)]?\s*[a-zA-Z0-9-]+)+)")
JUNK_EMAIL = re.compile(
    r"\.(png|jpe?g|gif|svg|webp|css|js)$|sentry|wixpress|example\.(com|org)|"
    r"yourcompany|yourdomain|domain\.com|email\.com$|^user@|^name@|^you@", re.I)
# policy-generator vendors whose own address often appears in template policies
VENDOR_DOMAINS = {"termly.io", "iubenda.com", "freeprivacypolicy.com", "privacypolicies.com",
                  "privacypolicygenerator.info", "termsfeed.com", "websitepolicies.com",
                  "getterms.io", "sentry.io"}
# hosts where "the site's contact page" is the platform's, not the developer's
GENERIC_HOSTS = {"docs.google.com", "sites.google.com", "drive.google.com", "notion.site",
                 "notion.so", "github.com", "gist.github.com", "medium.com", "termly.io",
                 "iubenda.com", "app.termly.io", "freeprivacypolicy.com", "dropbox.com",
                 "apps.apple.com", "play.google.com", "firebaseapp.com", "web.app"}
PRIVACY_WORDS = ("privacy", "dpo", "gdpr", "data", "legal", "dataprotection")
NO_RETRY_FETCH = ("SSLError", "TooManyRedirects", "InvalidURL", "MissingSchema", "InvalidSchema")
BROWSER_STATUSES = {"no_contact_found", "http_403", "http_406", "http_412", "http_429"}


# ------------------------------------------------------------------ helpers ---
def expired():
    return DEADLINE is not None and time.time() > DEADLINE


def safe_host(url):
    try:
        h = urlparse(url).netloc.lower()
    except ValueError:
        return ""
    return h[4:] if h.startswith("www.") else h


def is_transient(status):
    """Failures worth retrying with backoff."""
    if not isinstance(status, str):
        return False
    if status in ("http_429", "http_408") or status.startswith("http_5") or status.startswith("error_"):
        return True
    if status.startswith("fetch_failed_"):
        return not any(x in status for x in NO_RETRY_FETCH)
    return False


def needs_browser(status):
    return isinstance(status, str) and (status in BROWSER_STATUSES or status.startswith("browser_failed_"))


# ----------------------------------------------------------------- cleaning ---
def to_bool(v):
    return str(v).strip().lower() in {"true", "yes", "y", "1", "1.0"}


def parse_downloads(v):
    """'10,000,000+' -> 10000000 ; '5M+' -> 5000000 ; 1200 -> 1200"""
    if pd.isna(v):
        return 0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().upper().replace(",", "").replace("+", "").replace(" ", "")
    mult = {"K": 1e3, "M": 1e6, "B": 1e9}.get(s[-1:], 1)
    if mult != 1:
        s = s[:-1]
    try:
        return float(s) * mult
    except ValueError:
        return 0


def clean_and_dedupe(df):
    df.columns = [str(c).strip() for c in df.columns]
    required = ["Developer ID", "Rating", "Downloads", "Contains Ads", "Genre",
                "Category", "App Name", "App Link", "Privacy Policy"]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise SystemExit(f"Missing columns in input file: {missing}")

    total = len(df)
    df = df[df["Contains Ads"].apply(to_bool)].copy()
    print(f"[1] Kept {len(df):,} of {total:,} apps that contain ads")

    if "Developer Name" in df.columns:
        df["Developer ID"] = df["Developer ID"].fillna(df["Developer Name"])
    df = df.dropna(subset=["Developer ID"])

    df["_downloads"] = df["Downloads"].apply(parse_downloads)
    df["Rating"] = pd.to_numeric(df["Rating"], errors="coerce")

    agg = df.groupby("Developer ID").agg(
        Avg_Rating=("Rating", "mean"),
        Total_Downloads=("_downloads", "sum"),
        Apps_Count=("App Name", "count"),
    )
    top = df.sort_values("_downloads", ascending=False).drop_duplicates("Developer ID", keep="first")
    top = top.set_index("Developer ID")

    out = top.drop(columns=["Rating", "Downloads", "_downloads"]).join(agg)
    out["Avg_Rating"] = out["Avg_Rating"].round(2)
    out["Total_Downloads"] = out["Total_Downloads"].astype("int64")
    out = out.rename(columns={"Avg_Rating": "Avg Rating",
                              "Total_Downloads": "Total Downloads",
                              "Apps_Count": "Apps Count"})
    out = out.reset_index()
    out = out.sort_values("Total Downloads", ascending=False).reset_index(drop=True)
    print(f"[2] {len(out):,} unique developers after deduplication")
    return out


# ------------------------------------------------------ contact extraction ---
def region_from_country(country):
    c = str(country or "").strip().upper()
    return c if len(c) == 2 and c.isalpha() else None


def rank_email(email, page_host, dev_email_domain):
    local, _, domain = email.partition("@")
    score = 0
    if any(w in local for w in PRIVACY_WORDS):
        score += 100
    if dev_email_domain and domain == dev_email_domain:
        score += 20
    if page_host and (domain in page_host or page_host.endswith(domain)):
        score += 10
    if local in {"noreply", "no-reply", "donotreply"}:
        score -= 100
    return -score


def extract_contacts(html, page_url, region, dev_email):
    soup = BeautifulSoup(html, "html.parser")
    emails, phones = set(), []

    for a in soup.select('a[href^="mailto:"]'):
        # a mailto can hold several addresses: mailto:a@x.com,b@y.com
        for part in re.split(r"[,;]", a["href"][7:].split("?")[0]):
            emails.add(part.strip().lower())
    for a in soup.select('a[href^="tel:"]'):
        try:
            n = phonenumbers.parse(a["href"][4:], region)
            if phonenumbers.is_valid_number(n):
                phones.append(phonenumbers.format_number(n, phonenumbers.PhoneNumberFormat.INTERNATIONAL))
        except phonenumbers.NumberParseException:
            pass

    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    text = soup.get_text(" ")

    emails |= {e.lower() for e in EMAIL_RE.findall(text)}
    for user, host in OBFUSCATED_RE.findall(text):
        host = re.sub(r"\s*[\[\(]?\s*(?:dot|DOT)\s*[\]\)]?\s*", ".", host).replace(" ", "")
        emails.add(f"{user}@{host}".lower())

    emails = {e.strip(".,;:") for e in emails}
    emails = {e for e in emails
              if EMAIL_RE.fullmatch(e) and not JUNK_EMAIL.search(e)
              and e.split("@")[1] not in VENDOR_DOMAINS}

    for m in phonenumbers.PhoneNumberMatcher(text, region, leniency=phonenumbers.Leniency.VALID):
        phones.append(phonenumbers.format_number(m.number, phonenumbers.PhoneNumberFormat.INTERNATIONAL))

    page_host = safe_host(page_url)
    dev_domain = dev_email.split("@")[1].lower() if isinstance(dev_email, str) and "@" in dev_email else None
    ranked = sorted(emails, key=lambda e: rank_email(e, page_host, dev_domain))
    phones = list(dict.fromkeys(phones))
    return ranked[:MAX_EMAILS], phones[:MAX_PHONES]


def find_contact_links(html, page_url, limit=2):
    """Same-site links that look like a Contact page (best first), plus common guesses."""
    host = safe_host(page_url)
    soup = BeautifulSoup(html, "html.parser")
    scored = {}
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        if href.lower().startswith(("mailto:", "tel:", "javascript:", "#")):
            continue
        link = urljoin(page_url, href).split("#")[0]
        if not link.lower().startswith("http") or safe_host(link) != host or link == page_url:
            continue
        text = a.get_text(" ", strip=True).lower()
        low = link.lower()
        score = 0
        if "contact" in text:
            score += 3
        if "contact" in low:
            score += 2
        if score == 0 and ("get in touch" in text or "support" in text):
            score = 1
        if score:
            scored[link] = max(score, scored.get(link, 0))
    ranked = [l for l, _ in sorted(scored.items(), key=lambda kv: -kv[1])]
    if not ranked:
        p = urlparse(page_url)
        origin = f"{p.scheme}://{p.netloc}"
        ranked = [origin + "/contact", origin + "/contact-us"]
    return ranked[:limit]


# ---------------------------------------------------- HTTP fetch (threads) ---
_host_lock = threading.Lock()
_host_last = {}


def polite_wait(host):
    with _host_lock:
        now = time.time()
        wait = _host_last.get(host, 0) + HOST_DELAY - now
        _host_last[host] = max(now, _host_last.get(host, 0) + HOST_DELAY)
    if wait > 0:
        time.sleep(wait)


def _retry_after(resp):
    try:
        return min(60, int(resp.headers.get("Retry-After", "")))
    except ValueError:
        return None


def fetch(url, retries=2):
    """Returns (html_or_None, status). Retries 429 / 5xx / timeouts / connection errors
    with exponential backoff (honouring Retry-After)."""
    host = safe_host(url)
    last = "fetch_failed_unknown"
    for attempt in range(retries + 1):
        if expired():
            return None, "skipped_deadline"
        polite_wait(host)
        delay = None
        try:
            with requests.get(url, headers=HEADERS, timeout=TIMEOUT, stream=True) as r:
                code = r.status_code
                if code >= 400:
                    last = f"http_{code}"
                    if code in (408, 429) or code >= 500:
                        delay = _retry_after(r)
                    else:
                        return None, last
                else:
                    ctype = r.headers.get("content-type", "").lower()
                    if "pdf" in ctype or url.lower().endswith(".pdf"):
                        return None, "pdf_skipped"
                    raw = r.raw.read(MAX_BYTES, decode_content=True)
                    try:
                        return raw.decode(r.encoding or "utf-8", errors="ignore"), "ok"
                    except LookupError:
                        return raw.decode("utf-8", errors="ignore"), "ok"
        except requests.RequestException as e:
            name = type(e).__name__
            last = f"fetch_failed_{name}"
            if any(x in name for x in NO_RETRY_FETCH):
                return None, last
        if attempt < retries:
            time.sleep(delay if delay else min(60, 2 ** (attempt + 1) * 2) + random.random())
    return None, last


def scrape_one(url, country, dev_email, retries):
    if expired():
        return None
    html, status = fetch(url, retries)
    if status == "skipped_deadline":
        return None
    if html is None:
        return {"privacy_email": "", "privacy_number": "", "privacy_status": status,
                "privacy_source_url": ""}
    emails, phones = extract_contacts(html, url, region_from_country(country), dev_email)
    return {
        "privacy_email": "; ".join(emails),
        "privacy_number": "; ".join(phones),
        "privacy_status": "ok" if (emails or phones) else "no_contact_found",
        "privacy_source_url": url if (emails or phones) else "",
    }


def run_http_pass(jobs, cache, cache_path, workers, retries):
    done, shown = 0, 0
    with ThreadPoolExecutor(max_workers=workers) as ex:
        futures = {ex.submit(scrape_one, u, c, e, retries): u for u, (c, e) in jobs.items()}
        for fut in as_completed(futures):
            url = futures[fut]
            try:
                res = fut.result()
            except Exception as e:                      # keep the real error message
                if shown < 5:
                    shown += 1
                    print(f"    CRASH on {url}\n{traceback.format_exc()}")
                res = {"privacy_email": "", "privacy_number": "",
                       "privacy_status": f"error_{type(e).__name__}: {str(e)[:100]}",
                       "privacy_source_url": ""}
            if res is not None:
                cache[url] = res
            done += 1
            if done % 200 == 0:
                save_cache(cache_path, cache)
                print(f"    {done:,}/{len(jobs):,} done")
    save_cache(cache_path, cache)


# ------------------------------------------------- browser fallback (async) ---
async def _render(page, url, throttle):
    await throttle(safe_host(url))
    resp = await page.goto(url, wait_until="domcontentloaded", timeout=30000)
    try:
        await page.wait_for_load_state("networkidle", timeout=8000)
    except Exception:
        pass
    return (resp.status if resp else None), await page.content()


async def _browser_job(ctx, url, country, dev_email, throttle):
    page = await ctx.new_page()
    region = region_from_country(country)
    try:
        status, html = await _render(page, url, throttle)
        emails, phones = extract_contacts(html, url, region, dev_email)
        source = url
        if not emails and not phones and safe_host(url) not in GENERIC_HOSTS:
            for link in find_contact_links(html, url):
                try:
                    _, h2 = await _render(page, link, throttle)
                except Exception:
                    continue
                e2, p2 = extract_contacts(h2, link, region, dev_email)
                if e2 or p2:
                    emails, phones, source = e2, p2, link
                    break
        if emails or phones:
            st = "ok_browser"
        elif status and status >= 400:
            st = f"browser_http_{status}"
        else:
            st = "no_contact_found_browser"
        return {"privacy_email": "; ".join(emails), "privacy_number": "; ".join(phones),
                "privacy_status": st, "privacy_source_url": source if (emails or phones) else ""}
    except Exception as e:
        return {"privacy_email": "", "privacy_number": "",
                "privacy_status": f"browser_failed_{type(e).__name__}", "privacy_source_url": ""}
    finally:
        await page.close()


async def _browser_main(jobs, cache, cache_path, workers):
    from playwright.async_api import async_playwright

    host_next, host_lock = {}, asyncio.Lock()

    async def throttle(host):
        async with host_lock:
            now = time.time()
            at = max(now, host_next.get(host, 0))
            host_next[host] = at + HOST_DELAY
        if at > now:
            await asyncio.sleep(at - now)

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True, args=["--no-sandbox"])
        ctx = await browser.new_context(user_agent=BROWSER_UA, locale="en-US")

        async def block_heavy(route):
            if route.request.resource_type in ("image", "media", "font"):
                await route.abort()
            else:
                await route.continue_()
        await ctx.route("**/*", block_heavy)

        sem = asyncio.Semaphore(workers)
        done = 0

        async def run(url, country, dev_email):
            nonlocal done
            async with sem:
                if expired():
                    return
                try:
                    res = await asyncio.wait_for(_browser_job(ctx, url, country, dev_email, throttle), 120)
                except asyncio.TimeoutError:
                    res = {"privacy_email": "", "privacy_number": "",
                           "privacy_status": "browser_failed_Timeout", "privacy_source_url": ""}
                cache[url] = res
                done += 1
                if done % 100 == 0:
                    save_cache(cache_path, cache)
                    print(f"    browser {done:,}/{len(jobs):,} done")

        await asyncio.gather(*(run(u, c, e) for u, (c, e) in jobs.items()))
        await browser.close()


def run_browser_pass(jobs, cache, cache_path, workers):
    if not jobs:
        return
    try:
        import playwright.async_api  # noqa: F401
    except ImportError:
        print("    Playwright is not installed: pip install playwright && "
              "python -m playwright install chromium")
        return
    try:
        asyncio.run(_browser_main(jobs, cache, cache_path, workers))
    except Exception as e:
        print(f"    Browser pass aborted: {type(e).__name__}: {str(e)[:200]}\n"
              f"    (if the browser is missing run: python -m playwright install chromium)")
    save_cache(cache_path, cache)


# -------------------------------------------------------------------- cache ---
def load_cache(path):
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_cache(path, cache):
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(cache, f)
    os.replace(tmp, path)


# ---------------------------------------------------------------- enrichment ---
def enrich_privacy(df, cache_path, args):
    cache = load_cache(cache_path)

    url_info = {}
    for _, row in df.iterrows():
        u = row["Privacy Policy"]
        if isinstance(u, str) and u.strip().lower().startswith("http"):
            url_info.setdefault(u.strip(), (row.get("Country"), row.get("Developer Email")))

    if args.retry_failed:
        drop = [u for u in url_info if u in cache and is_transient(cache[u].get("privacy_status"))]
        for u in drop:
            del cache[u]
        print(f"[3a] Retry pass: {len(drop):,} transient failures queued (rows that worked are untouched)")

    jobs = {u: i for u, i in url_info.items() if u not in cache}
    print(f"[3] HTTP pass: {len(jobs):,} pages to fetch ({len(url_info) - len(jobs):,} already cached)")
    if jobs:
        run_http_pass(jobs, cache, cache_path, args.workers, args.retries)

    if args.browser_fallback and not expired():
        bjobs = {u: i for u, i in url_info.items()
                 if u in cache and needs_browser(cache[u].get("privacy_status"))}
        print(f"[3b] Browser pass: {len(bjobs):,} pages (empty / 403 / 429 rows only)")
        run_browser_pass(bjobs, cache, cache_path, args.browser_workers)

    no_url = {"privacy_email": "", "privacy_number": "", "privacy_status": "no_privacy_url",
              "privacy_source_url": ""}
    pending = {"privacy_email": "", "privacy_number": "", "privacy_status": "not_scraped_yet",
               "privacy_source_url": ""}

    def lookup(u):
        if not (isinstance(u, str) and u.strip().lower().startswith("http")):
            return no_url
        return cache.get(u.strip(), pending)

    res = df["Privacy Policy"].apply(lookup)
    emails, phones, statuses, sources = [], [], [], []
    for d in res:
        # drop policy-generator vendor addresses that older cached runs may still hold
        kept = [e for e in d.get("privacy_email", "").split("; ")
                if e and e.split("@")[-1] not in VENDOR_DOMAINS]
        st = d.get("privacy_status", "")
        if st.startswith("ok") and not kept and not d.get("privacy_number"):
            st = "vendor_email_only"
        emails.append("; ".join(kept))
        phones.append(d.get("privacy_number", ""))
        statuses.append(st)
        sources.append(d.get("privacy_source_url", ""))
    df["privacy_email"], df["privacy_number"] = emails, phones
    df["privacy_status"], df["privacy_source_url"] = statuses, sources
    return df


# -------------------------------------------------------------------- output ---
def write_excel(df, path):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, index=False, sheet_name="Developers")
        ws = xw.sheets["Developers"]
        head_fill = PatternFill("solid", fgColor="1F3864")
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial", size=10)
        for i, col in enumerate(df.columns, start=1):
            width = min(max(len(str(col)), int(df[col].astype(str).str.len().quantile(0.9))) + 2, 50)
            ws.column_dimensions[get_column_letter(i)].width = max(width, 12)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        if "Total Downloads" in df.columns:
            col = get_column_letter(list(df.columns).index("Total Downloads") + 1)
            for c in ws[col][1:]:
                c.number_format = "#,##0"
    print(f"[4] Saved {path}")


def main():
    global HOST_DELAY, DEADLINE
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--workers", type=int, default=12)
    ap.add_argument("--limit", type=int, default=0, help="only process the top N developers (testing)")
    ap.add_argument("--no-scrape", action="store_true")
    ap.add_argument("--retry-failed", action="store_true",
                    help="re-fetch transient failures (429, 5xx, timeouts, crashes) with backoff")
    ap.add_argument("--browser-fallback", action="store_true",
                    help="headless-browser retry for no_contact_found / 403 rows, follows Contact links")
    ap.add_argument("--browser-workers", type=int, default=4)
    ap.add_argument("--retries", type=int, default=2, help="backoff retries per page")
    ap.add_argument("--host-delay", type=float, default=1.0, help="seconds between hits to one host")
    ap.add_argument("--max-minutes", type=float, default=0, help="stop gracefully after N minutes")
    args = ap.parse_args()

    HOST_DELAY = args.host_delay
    if args.max_minutes:
        DEADLINE = time.time() + args.max_minutes * 60

    df = pd.read_excel(args.input)
    df = clean_and_dedupe(df)
    if args.limit:
        df = df.head(args.limit)
    if not args.no_scrape:
        df = enrich_privacy(df, args.output + ".cache.json", args)
        print(df["privacy_status"].str.split(":").str[0].value_counts().to_string())
    write_excel(df, args.output)
    if expired():
        print("Stopped at the time limit - re-run the same command to continue.")
        sys.exit(3)


if __name__ == "__main__":
    main()
